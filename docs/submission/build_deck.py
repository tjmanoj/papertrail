"""Builds the PaperTrail submission deck inside the official hackathon template.

Keeps the template's branding (black header band, gradient footer, Thank You
slide) and fills the blank content slides. Arial throughout, because anything
else falls back badly on stock macOS.
"""
import copy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = "template.pptx"
OUT = "PaperTrail-Provenance-Prototype-Submission.pptx"

# geometry (EMU) ---------------------------------------------------------
SLIDE_W, SLIDE_H = 9144000, 5143500
MARGIN = 430000
BODY_W = SLIDE_W - 2 * MARGIN
TOP = 620000                      # below the black header band
BOTTOM = 4880000                  # above the gradient footer

INK = RGBColor(0x10, 0x1A, 0x2B)
INK2 = RGBColor(0x45, 0x55, 0x6E)
MUTED = RGBColor(0x79, 0x8A, 0xA2)
CYAN = RGBColor(0x0E, 0x88, 0xB4)
CYAN_BRIGHT = RGBColor(0x29, 0xB5, 0xE8)
OK = RGBColor(0x1A, 0x6F, 0x45)
CRIT = RGBColor(0x9E, 0x27, 0x27)
PANEL = RGBColor(0xF2, 0xF5, 0xF9)
LINE = RGBColor(0xD6, 0xDE, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation(SRC)


def set_text(tf, runs, *, size=14, color=INK, bold=False, space_after=4,
             align=PP_ALIGN.LEFT, line=None):
    """runs: list of paragraphs; each is a str or list of (text, overrides)."""
    tf.word_wrap = True
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        if line:
            p.line_spacing = line
        pieces = [(para, {})] if isinstance(para, str) else para
        for text, ov in pieces:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = "Arial"
            f.size = Pt(ov.get("size", size))
            f.bold = ov.get("bold", bold)
            f.color.rgb = ov.get("color", color)


def box(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(Emu(l), Emu(t), Emu(w), Emu(h))
    tb.text_frame.word_wrap = True
    return tb.text_frame


def panel(slide, l, t, w, h, fill=PANEL, line_col=LINE):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(l), Emu(t), Emu(w), Emu(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line_col
    sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = 0.06
    except Exception:
        pass
    sh.text_frame.word_wrap = True
    sh.text_frame.margin_left = Emu(130000)
    sh.text_frame.margin_right = Emu(130000)
    sh.text_frame.margin_top = Emu(90000)
    sh.text_frame.margin_bottom = Emu(90000)
    return sh


def title(slide, text, sub=None):
    tf = box(slide, MARGIN, TOP - 60000, BODY_W, 520000)
    set_text(tf, [text], size=26, bold=True, color=INK, space_after=2)
    if sub:
        tf2 = box(slide, MARGIN, TOP + 330000, BODY_W, 300000)
        set_text(tf2, [sub], size=12.5, color=MUTED)
    return TOP + (690000 if sub else 520000)


def extract_background(source, path="_bg.png"):
    """Save the template's content-slide background so it can be re-added cleanly.
    Deep-copying the picture element instead would carry a relationship id that
    does not exist on the new slide, which corrupts the file."""
    for sh in source.shapes:
        if sh.shape_type == 13:  # PICTURE
            with open(path, "wb") as fh:
                fh.write(sh.image.blob)
            return path
    return None


def add_background(target, path):
    pic = target.shapes.add_picture(path, Emu(0), Emu(0), Emu(SLIDE_W), Emu(SLIDE_H))
    target.shapes._spTree.remove(pic._element)
    target.shapes._spTree.insert(2, pic._element)   # send to back


# ---------------------------------------------------------------- slide 1
s1 = prs.slides[0]
VALUES = {
    "Team Name :": "Team Name :  Provenance",
    "Team Leader Name :": "Team Leader Name :  Manoj T",
    "Team Size :": "Team Size :  2",
    "Problem Statement :": "Problem Statement :  1 — Risk, Fraud and Regulatory Intelligence Copilot",
}
for sh in s1.shapes:
    if not sh.has_text_frame:
        continue
    key = sh.text_frame.text.strip()
    if key in VALUES:
        tf = sh.text_frame
        tf.clear()
        label, _, val = VALUES[key].partition(":")
        set_text(tf, [[(label + ": ", {"color": INK2, "bold": False}),
                       (val.strip(), {"color": INK, "bold": True})]],
                 size=15)

# Use the template's slides exactly as they are: no adds, no removes, no
# reordering of sldIdLst - that corrupts the package. Slide 2 carries the
# submission guidelines, which are scaffolding, so its text box is emptied and
# the slide reused for content.
for sh in list(prs.slides[1].shapes):
    if sh.has_text_frame and "Submission Guidelines" in sh.text_frame.text:
        sh._element.getparent().remove(sh._element)

S = [prs.slides[i] for i in range(1, 5)]   # four content slides, 6th is Thank You

# ---------------------------------------------------------------- 2. problem
sl = S[0]
y = title(sl, "PaperTrail — the problem",
          "Risk, Fraud and Regulatory Intelligence Copilot · banking and NBFC")

p = panel(sl, MARGIN, y, BODY_W, 900000, fill=RGBColor(0x0E, 0x16, 0x24), line_col=RGBColor(0x0E, 0x16, 0x24))
set_text(p.text_frame,
         [[("Finding the risk signal is the cheap part. ", {"color": WHITE}),
           ("Defending the number is the expensive part.", {"color": CYAN_BRIGHT, "bold": True})]],
         size=17, align=PP_ALIGN.CENTER)
y += 1050000

cols = [
    ("THE PAIN TODAY",
     "An analyst answers “which transactions breached our AML thresholds last month” in ten "
     "minutes — then spends two days assembling an evidence pack an auditor will accept.\n\n"
     "Dashboards hand you numbers nobody can defend. Two analysts asked the same question "
     "compute it differently: booking or value date, reversals in or out."),
    ("THE USER",
     "A compliance analyst or MLRO at a mid-size commercial bank — the person who has to put "
     "their name on the number.\n\n"
     "Banking is India’s largest GCC segment; regulatory reporting and AML operations are the "
     "day job of tens of thousands of those staff."),
]
cw = (BODY_W - 240000) // 2
for i, (head, body) in enumerate(cols):
    tf = box(sl, MARGIN + i * (cw + 240000), y, cw, 1500000)
    set_text(tf, [head], size=10, bold=True, color=CYAN, space_after=6)
    set_text_body = tf.add_paragraph()
    set_text_body.space_after = Pt(0)
    for line_ in body.split("\n\n"):
        pp = tf.add_paragraph()
        pp.space_after = Pt(7)
        r = pp.add_run(); r.text = line_
        r.font.name = "Arial"; r.font.size = Pt(10.5); r.font.color.rgb = INK2

y += 1620000
p2 = panel(sl, MARGIN, y, BODY_W, 700000, fill=PANEL)
set_text(p2.text_frame,
         [[("What PaperTrail produces is not a chat reply. It is a filing — and every figure in it "
            "is a footnote that resolves to its governed metric definition, the exact SQL, the "
            "source rows, and the regulation clause that makes it matter.", {})]],
         size=12.5, color=INK)

# ---------------------------------------------------------------- 3. architecture
sl = S[1]
y = title(sl, "Architecture", "Snowflake-native end to end · no LangChain, no external vector store, no third-party LLM API")

ROW_H = 560000
GAP = 150000
half = (BODY_W - GAP) // 2

def stack(x, w, rows, head_col=CYAN):
    yy = y
    for i, (head, body) in enumerate(rows):
        sh = panel(sl, x, yy, w, ROW_H)
        tf = sh.text_frame
        set_text(tf, [head], size=10, bold=True, color=head_col, space_after=2)
        pp = tf.add_paragraph(); pp.space_after = Pt(0)
        r = pp.add_run(); r.text = body
        r.font.name = "Arial"; r.font.size = Pt(10); r.font.color.rgb = INK2
        yy += ROW_H + GAP
    return yy

left_rows = [
    ("STRUCTURED  ·  synthetic banking data",
     "240,798 transactions · 809 counterparties · 1,497 accounts · 18 months"),
    ("RAW → CURATED → GOLD   ·  9 Dynamic Tables",
     "Entity resolution; governance choices kept as explicit columns, not filtered away"),
    ("SEMANTIC VIEW   ·  Cortex Analyst",
     "7 governed metrics · 10 verified queries · each metric states its governance decision"),
]
right_rows = [
    ("UNSTRUCTURED  ·  11 regulatory PDFs",
     "Parsed inside Snowflake with AI_PARSE_DOCUMENT → 72 clauses"),
    ("CORTEX SEARCH   ·  REGULATION_SEARCH",
     "Clause retrieval with citation attributes · 11/12 top-1, 12/12 top-3"),
    ("ROW ACCESS POLICIES",
     "AU / SG analyst personas, justified against AUSTRAC Pt 8.1 and MAS Notice 626 §13"),
]
yl = stack(MARGIN, half, left_rows)
stack(MARGIN + half + GAP, half, right_rows)

agent = panel(sl, MARGIN, yl, BODY_W, 480000,
              fill=RGBColor(0x0E, 0x16, 0x24), line_col=RGBColor(0x0E, 0x16, 0x24))
set_text(agent.text_frame,
         [[("CORTEX AGENT   ", {"color": CYAN_BRIGHT, "bold": True, "size": 11}),
           ("figures → Analyst · rules → Search · both when a number needs a rule · "
            "refuses what it cannot ground", {"color": WHITE, "size": 11})]],
         align=PP_ALIGN.CENTER)
yl += 480000 + GAP

# skills chain
skills = [
    ("risk-scanner", "governed query +\nprovenance bundle"),
    ("regulation-linker", "clause retrieval\nwith citations"),
    ("finding-writer", "assembles finding;\ncannot invent a figure"),
    ("provenance-logger", "append-only,\ncontent-hashed"),
]
sw = (BODY_W - 3 * 110000) // 4
for i, (name, desc) in enumerate(skills):
    sh = panel(sl, MARGIN + i * (sw + 110000), yl, sw, 540000,
               fill=WHITE, line_col=CYAN)
    tf = sh.text_frame
    set_text(tf, [name], size=10.5, bold=True, color=CYAN, space_after=2, align=PP_ALIGN.CENTER)
    pp = tf.add_paragraph(); pp.alignment = PP_ALIGN.CENTER; pp.space_after = Pt(0)
    r = pp.add_run(); r.text = desc
    r.font.name = "Arial"; r.font.size = Pt(9); r.font.color.rgb = INK2
    if i < 3:
        ar = box(sl, MARGIN + i * (sw + 110000) + sw + 14000, yl + 180000, 90000, 200000)
        set_text(ar, ["→"], size=13, color=CYAN, align=PP_ALIGN.CENTER)

tf = box(sl, MARGIN, yl + 580000, BODY_W, 300000)
set_text(tf, ["Four reusable CoCo skills, each with a documented interface and an explicit boundary. "
              "finding-writer cannot state an ungrounded figure: its prompt contains only extracted facts "
              "with no database access, and every number in its output is validated back to an input fact."],
         size=10, color=MUTED)

# ---------------------------------------------------------------- 4. proof
sl = S[2]
y = title(sl, "We measured the claim everyone else asserts",
          "Governed vs ungoverned — same data, same model, full raw DDL, no hint of any governance rule")

p = panel(sl, MARGIN, y, BODY_W, 640000, fill=RGBColor(0xFA, 0xE8, 0xE8), line_col=RGBColor(0xE4, 0xBF, 0xBF))
set_text(p.text_frame,
         [[("Our hypothesis was wrong. ", {"bold": True, "color": CRIT}),
           ("An ungoverned model is not inconsistent — across 25 runs the SQL was byte-identical. "
            "It is perfectly confident and materially wrong on 2 of 5 questions.", {"color": INK})]],
         size=12.5)
y += 760000

rows = [
    ("Suspicious volume, Q3 2026", "8,797,923", "4,295,229", "+105%", CRIT),
    ("Exposure to CP-STRUCT-01", "87,185", "87,185", "agree*", MUTED),
    ("Alert closure rate 2026", "63.6%", "63.6%", "agree", MUTED),
    ("SAR filing share", "6.5%", "6.5%", "agree", MUTED),
    ("High-risk counterparties, Singapore", "23", "36", "−36%", CRIT),
]
colw = [BODY_W - 2400000 - 900000, 1200000, 1200000, 900000]
hy = y
hdrs = ["QUESTION", "UNGOVERNED", "GOVERNED", ""]
xx = MARGIN
for i, h in enumerate(hdrs):
    tf = box(sl, xx, hy, colw[i], 220000)
    set_text(tf, [h], size=8.5, bold=True, color=MUTED,
             align=PP_ALIGN.RIGHT if i else PP_ALIGN.LEFT)
    xx += colw[i]
hy += 250000
for label, ung, gov, verdict, col in rows:
    xx = MARGIN
    for i, (val, al) in enumerate([(label, PP_ALIGN.LEFT), (ung, PP_ALIGN.RIGHT),
                                   (gov, PP_ALIGN.RIGHT), (verdict, PP_ALIGN.RIGHT)]):
        tf = box(sl, xx, hy, colw[i], 260000)
        set_text(tf, [val], size=11,
                 color=col if i == 3 else INK,
                 bold=(i == 3 and col is CRIT), align=al)
        xx += colw[i]
    hy += 285000

tf = box(sl, MARGIN, hy + 40000, BODY_W, 560000)
set_text(tf, [
    [("* Q2 agrees by coincidence", {"bold": True, "color": INK, "size": 10.5}),
     (" — it computes direct exposure rather than entity-resolved, and matches only because that "
      "counterparty has no beneficial-ownership links. Agreement is not evidence of a sound method.",
      {"color": MUTED, "size": 10.5})],
    [("The $4.5M overstatement comes from four missed decisions at once: booking date rather than value "
      "date, reversals counted, unsettled counted, and transactions on already-closed alerts treated as "
      "suspicious. Nothing in the output signals any of it.", {"color": INK2, "size": 10.5})],
], space_after=5)

# ---------------------------------------------------------------- 5. impact
sl = S[3]
y = title(sl, "Measured outcomes", "Scored against an answer key that was never loaded into Snowflake")

tiles = [
    ("0.944", "detection recall\n17 of 18 planted positives"),
    ("5 of 6", "typologies at F1 1.000"),
    ("11/12", "clause retrieval, top-1\n12/12 at top-3"),
    ("5/5", "agent tool-routing\nincluding the refusal"),
]
tw = (BODY_W - 3 * 120000) // 4
for i, (v, l) in enumerate(tiles):
    sh = panel(sl, MARGIN + i * (tw + 120000), y, tw, 760000, fill=WHITE, line_col=LINE)
    tf = sh.text_frame
    set_text(tf, [v], size=22, bold=True, color=CYAN, space_after=2, align=PP_ALIGN.CENTER)
    pp = tf.add_paragraph(); pp.alignment = PP_ALIGN.CENTER; pp.space_after = Pt(0)
    r = pp.add_run(); r.text = l
    r.font.name = "Arial"; r.font.size = Pt(8.5); r.font.color.rgb = MUTED
y += 880000

left = panel(sl, MARGIN, y, half, 1180000, fill=PANEL)
set_text(left.text_frame, [
    [("TIME SAVED, HONESTLY STATED", {"size": 10, "bold": True, "color": CYAN})],
    "The evidence pack a finding produces — figure, metric definition, SQL, source rows, "
    "clause — is the artefact that takes an analyst two days to assemble by hand. "
    "PaperTrail emits it with the answer, and the figure re-derives from its own footnote "
    "in about 300 ms.",
], size=10.5, color=INK2)

right = panel(sl, MARGIN + half + GAP, y, half, 1180000, fill=PANEL)
set_text(right.text_frame, [
    [("WHY THIS IS NOT 100%", {"size": 10, "bold": True, "color": CYAN})],
    "Velocity-spike detection is weak — 20 false positives. Removing it would take overall "
    "precision and F1 to 1.000. We kept it: velocity monitoring is a regulatory requirement, "
    "not an optional enhancement. A perfect score on self-authored data, detectors and answer "
    "key would be evidence of tuning, not detection.",
], size=10.5, color=INK2)
y += 1300000

tf = box(sl, MARGIN, y, BODY_W, 700000)
set_text(tf, [
    [("Scalability and beyond the demo.  ", {"bold": True, "color": INK, "size": 11}),
     ("Every object is idempotent DDL and every row is generated, so the system rebuilds "
      "byte-for-byte in any Snowflake account. The four skills are parameterised for reuse: swap "
      "the semantic view and the clause corpus, keep the provenance machinery.",
      {"color": INK2, "size": 10.5})],
], line=1.15)

# ------------------------------------------- links footer on the impact slide
sl = S[3]
p = panel(sl, MARGIN, 4120000, BODY_W, 760000,
          fill=RGBColor(0x0E, 0x16, 0x24), line_col=RGBColor(0x0E, 0x16, 0x24))
set_text(p.text_frame, [
    [("Live app   ", {"color": MUTED, "size": 9.5}),
     ("papertrail-provenance.vercel.app", {"color": CYAN_BRIGHT, "size": 12.5, "bold": True}),
     ("   no login required   ·   Repo   ", {"color": MUTED, "size": 9.5}),
     ("github.com/tjmanoj/papertrail", {"color": WHITE, "size": 10.5})],
    [("Native app   ", {"color": MUTED, "size": 9.5}),
     ("PAPERTRAIL.GOLD.PAPERTRAIL_APP  ·  Streamlit in Snowflake", {"color": WHITE, "size": 10})],
    [("CoCo evidence   ", {"color": MUTED, "size": 9.5}),
     ("docs/coco/ — 11 full session transcripts and 26 verbatim prompts, each re-derivable with "
      "cortex conversations transcript <id>", {"color": WHITE, "size": 9.5})],
], space_after=2)

# Final sweep: Arial on every run. Manrope survives from the template's own
# text boxes and is not installed on stock macOS, so it would fall back badly.
for _sl in prs.slides:
    for _sh in _sl.shapes:
        if not _sh.has_text_frame:
            continue
        for _p in _sh.text_frame.paragraphs:
            for _r in _p.runs:
                _r.font.name = "Arial"

prs.save(OUT)
print("saved:", OUT)
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
